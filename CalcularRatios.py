from openpyxl import load_workbook, Workbook
import os

ruta = os.path.join(
    os.path.expanduser("~"),
    "Documents",
    "Tablas",
    "SigncryptionBenchmarks500.xlsx"
)

wb = load_workbook(ruta)
ws = wb.active

weier = {}

ratios = []

for i,row in enumerate(ws.iter_rows(values_only=True)):

    if i == 0:
        continue

    curva, ecuacion, msg, p, exp, tc, td, tt = row

    key = (p,msg,exp)

    if curva == "Weierstrass":
        weier[key] = (tc,td,tt)

    elif curva == "Edwards":

        if key in weier:

            tc_w, td_w, tt_w = weier[key]

            ratio_c = tc / tc_w
            ratio_d = td / td_w
            ratio_t = tt / tt_w

            ratios.append((exp,p,msg,ratio_c,ratio_d,ratio_t))


wb2 = Workbook()
ws2 = wb2.active

ws2.append([
    "Experimento",
    "p",
    "MsgSize",
    "RatioCifrado",
    "RatioDescifrado",
    "RatioTotal"
])

for r in ratios:
    ws2.append(r)

salida = os.path.join(
    os.path.expanduser("~"),
    "Documents",
    "Tablas",
    "tabla_ratios500.xlsx"
)

wb2.save(salida)

print("Tabla de ratios guardada en:", salida)